"""Write the companies and their contact details to the outreach Excel tracker.

Same "Agency Tracker" layout, columns and styling as test_clutch/excel_writer.py,
with the contact columns (DM Name, DM Title, Email, LinkedIn, Contact Form Link)
filled from the website step, plus Phone and the pages each detail came from.
Scoring columns stay blank: the scoring module is not part of this project.
"""
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LEADS_XLSX_PATH = "output/clutch_outreach_tracker.xlsx"
LEADS_CSV_PATH = "output/clutch_companies.csv"
# CSV columns: every stored field. A field the listing or website did not have stays blank.
CSV_COLUMNS = [
    "agency_name", "country", "city", "website", "profile_link", "main_services",
    "secondary_services", "team_size_raw", "hourly_rate_raw", "min_project_size_raw", "rating",
    "review_count", "description", "emails", "phones", "decision_makers", "linkedin",
    "contact_form_url", "contact_source_urls", "enrichment_status", "enrichment_error",
    "source_url", "page_number",
]

OUTREACH_COLS = [
    "Date Added", "Platform", "Agency Name", "Country", "City", "Website",
    "Profile Link", "Main Services", "Secondary Services", "Team Size",
    "Hourly Rate (USD)", "Min Project Size (USD)", "Industries Served",
    "DM Name", "DM Title", "Email", "Phone", "LinkedIn Profile", "Contact Form Link",
    "Signs They May Outsource", "Qual Score", "Priority", "First Outreach Date",
    "Follow-Up 1 Date", "Follow-Up 2 Date", "Response Status", "Next Action", "Remarks",
]
DAILY_LOG_COLS = [
    "Log Date", "Researcher", "Researched", "Qualified", "DMs Found", "Emails Sent",
    "LinkedIn Sent", "Follow-Ups", "Replies", "Positive", "Calls Sched.", "Notes / Lessons",
]
CONTACT_DETAIL_COLS = ["All Emails", "All Phones", "Other Decision Makers", "Contact Sources", "Website Check"]

_NAVY = PatternFill("solid", fgColor="1F4E79")
_GREEN = PatternFill("solid", fgColor="375623")
_BLUE = PatternFill("solid", fgColor="2E75B6")
_WHITE = Font(color="FFFFFF", bold=True)
_WIDTHS = {
    "Date Added": 14, "Platform": 10, "Agency Name": 30, "Country": 15, "City": 18,
    "Website": 32, "Profile Link": 32, "Main Services": 42, "Secondary Services": 36,
    "Team Size": 12, "Hourly Rate (USD)": 18, "Min Project Size (USD)": 20,
    "Industries Served": 25, "DM Name": 22, "DM Title": 24, "Email": 30, "Phone": 20,
    "LinkedIn Profile": 34, "Contact Form Link": 34, "Signs They May Outsource": 45,
    "Qual Score": 13, "Priority": 18, "First Outreach Date": 18, "Follow-Up 1 Date": 18,
    "Follow-Up 2 Date": 18, "Response Status": 18, "Next Action": 25, "Remarks": 50,
    "Log Date": 14, "Researcher": 15, "Researched": 13, "Qualified": 13, "DMs Found": 12,
    "Emails Sent": 13, "LinkedIn Sent": 14, "Follow-Ups": 12, "Replies": 10, "Positive": 10,
    "Calls Sched.": 13, "Notes / Lessons": 30, "All Emails": 40, "All Phones": 30,
    "Other Decision Makers": 40, "Contact Sources": 50, "Website Check": 22,
}


def _apply_header(cell, text, fill):
    cell.value = text
    cell.fill = fill
    cell.font = _WHITE
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _split(value):
    return [part.strip() for part in (value or "").split(";") if part.strip()]


def _row(company, today):
    people = _split(company.get("decision_makers"))
    dm_name, _, dm_title = (people[0].partition(" (") if people else ("", "", ""))
    emails, phones = _split(company.get("emails")), _split(company.get("phones"))
    linkedin = _split(company.get("linkedin"))
    # The DM's own profile if its URL carries their name, otherwise the company page.
    name_parts = [part.lower() for part in dm_name.split() if len(part) > 2]
    person_links = [link for link in linkedin
                    if "/in/" in link and any(part in link.lower() for part in name_parts)]
    person_links += [link for link in linkedin if "/company/" in link]
    status = company.get("enrichment_status", "")
    if company.get("enrichment_error"):
        status += f" ({company['enrichment_error']})"
    return {
        "Date Added": today, "Platform": "Clutch",
        "Agency Name": company.get("agency_name", ""), "Country": company.get("country", ""),
        "City": company.get("city", ""), "Website": company.get("website", ""),
        "Profile Link": company.get("profile_link", ""),
        "Main Services": company.get("main_services", "").replace("; ", ", "),
        "Secondary Services": company.get("secondary_services", "").replace("; ", ", "),
        "Team Size": company.get("team_size_raw", ""), "Hourly Rate (USD)": company.get("hourly_rate_raw", ""),
        "Min Project Size (USD)": company.get("min_project_size_raw", ""),
        "DM Name": dm_name, "DM Title": dm_title.rstrip(")"),
        "Email": emails[0] if emails else "", "Phone": phones[0] if phones else "",
        "LinkedIn Profile": (person_links or linkedin or [""])[0],
        "Contact Form Link": company.get("contact_form_url", ""),
        "Response Status": "Not Contacted",
        "All Emails": "; ".join(emails), "All Phones": "; ".join(phones),
        "Other Decision Makers": "; ".join(people[1:]),
        "Contact Sources": company.get("contact_source_urls", ""), "Website Check": status,
    }


def write_csv(companies, output_path=LEADS_CSV_PATH):
    """All stored companies as CSV (UTF-8 with BOM so Excel opens it correctly)."""
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, restval="", extrasaction="ignore")
        writer.writeheader()
        for company in companies:
            writer.writerow({column: company.get(column) or "" for column in CSV_COLUMNS})
    return output_path


def write_tracker(companies, output_path=LEADS_XLSX_PATH):
    wb = Workbook()
    ws = wb.active
    ws.title = "Agency Tracker"
    groups = [("OUTREACH PIPELINE", OUTREACH_COLS, _NAVY), ("DAILY PROGRESS LOG", DAILY_LOG_COLS, _GREEN),
              ("CONTACT DETAILS (from company websites)", CONTACT_DETAIL_COLS, _BLUE)]
    all_cols, start = [], 1
    for title, cols, fill in groups:
        _apply_header(ws.cell(1, start), title, fill)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + len(cols) - 1)
        for offset, col in enumerate(cols):
            _apply_header(ws.cell(2, start + offset), col, fill)
        all_cols += cols
        start += len(cols)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36

    today = date.today().strftime("%d-%b-%Y")
    for company in companies:
        values = _row(company, today)
        ws.append([values.get(col, "") for col in all_cols])

    for index, col in enumerate(all_cols, 1):
        ws.column_dimensions[get_column_letter(index)].width = _WIDTHS.get(col, 15)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(all_cols))}{max(ws.max_row, 2)}"
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
