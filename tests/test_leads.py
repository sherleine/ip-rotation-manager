import csv
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from app.clutch_parser import build_page_url, extract_company_rows
from app.leads_db import LeadsDatabase
from app.leads_export import write_csv, write_tracker
from app.llm_extract import parse_people
from app import website_contacts
from app.webdata_client import AccessBlockedError
from app.website_contacts import (
    WebsiteContactFinder,
    decode_cfemail,
    extract_page_contacts,
    find_decision_makers,
    rank_emails,
    role_rank,
    unique_phones,
)

CARD = """
<div class="provider-list-item">
  <h3 class="provider__title">Example AI</h3>
  <a class="directory_profile" href="/profile/example-ai"></a>
  <div class="location">Mumbai, India</div>
  <div class="hourly-rate">$100 - $149 / hr</div>
  <div class="sg-rating__reviews">12 reviews</div>
</div>
"""
CONTACT_PAGE = """
<html><body>
  <h1>Contact us</h1>
  <a href="mailto:Hello@Acme-AI.com?subject=Hi">Email</a>
  <a href="/cdn-cgi/l/email-protection#{cf}">[email protected]</a>
  <a href="tel:+1 (555) 123-4567">Call</a>
  <p>Office: +44 20 7946 0958 - founded 2015 - logo@2x.png</p>
  <a href="https://www.linkedin.com/company/example-ai/?trk=x">LinkedIn</a>
  <h3>Jane Smith</h3><p>Founder &amp; CEO</p>
  <form><input type="email"><textarea></textarea></form>
</body></html>
"""


def cf_encode(email, key=0x42):
    return f"{key:02x}" + "".join(f"{ord(ch) ^ key:02x}" for ch in email)


class ClutchParserTests(unittest.TestCase):
    def test_page_url_and_card(self):
        self.assertEqual(build_page_url("https://example.test/list?x=1", 2), "https://example.test/list?x=1&page=2")
        rows = extract_company_rows(CARD, "https://clutch.co/developers/artificial-intelligence", 1)
        self.assertEqual(rows[0]["agency_name"], "Example AI")
        self.assertEqual(rows[0]["profile_link"], "https://clutch.co/profile/example-ai")
        self.assertEqual(rows[0]["review_count"], 12)


class WebsiteContactTests(unittest.TestCase):
    def test_cfemail_decoding(self):
        self.assertEqual(decode_cfemail(cf_encode("sales@example.com")), "sales@example.com")
        self.assertEqual(decode_cfemail("zz"), "")

    def test_contact_page_extraction(self):
        html = CONTACT_PAGE.format(cf=cf_encode("jane@acme-ai.com"))
        page = extract_page_contacts(BeautifulSoup(html, "lxml"), "https://acme-ai.com/contact")
        self.assertEqual(page["emails"], {"hello@acme-ai.com", "jane@acme-ai.com"})
        self.assertIn("+1 (555) 123-4567", page["phones"])
        self.assertIn("+44 20 7946 0958", page["phones"])
        self.assertFalse(any("2015" == phone for phone in page["phones"]))
        self.assertEqual(page["linkedin"], {"https://www.linkedin.com/company/example-ai"})
        self.assertEqual(page["decision_makers"], ["Jane Smith (Founder & CEO)"])
        self.assertTrue(page["contact_form"])

    def test_decision_maker_patterns(self):
        lines = ["Meet our team", "John O'Neil, Co-Founder & CTO", "Our Team", "Founder", "Read More"]
        self.assertEqual(find_decision_makers(lines), ["John O'Neil (Co-Founder & CTO)"])

    def test_client_testimonials_are_not_decision_makers(self):
        lines = ["Aleksei Turchak", "CTO", "John Sims", "CTO, STEMscopes by Accelerate Learning",
                 "Don Keeler", "Former CTO, Accelerate Learning Inc., USA", "Ann Lee", "CEO of Lumen Touch"]
        self.assertEqual(find_decision_makers(lines), ["Aleksei Turchak (CTO)"])

    def test_same_phone_in_two_formats_is_kept_once(self):
        phones = {"+1 (949) 620-0723", "+19496200723", "+37052075658"}
        self.assertEqual(unique_phones(phones), ["+1 (949) 620-0723", "+37052075658"])

    def test_owner_and_ceo_rank_before_cto(self):
        people = ["Aleksei Turchak (CTO)", "Sergey Tolkachev (CEO)", "Mia Ray (Co-Founder & COO)"]
        self.assertEqual(sorted(people, key=role_rank),
                         ["Mia Ray (Co-Founder & COO)", "Sergey Tolkachev (CEO)", "Aleksei Turchak (CTO)"])

    def test_phone_does_not_run_onto_the_next_line(self):
        soup = BeautifulSoup("<p>Call +1 415 286 3266</p><p>391 clients</p>", "lxml")
        self.assertEqual(extract_page_contacts(soup, "https://acme-ai.com")["phones"], {"+1 415 286 3266"})

    def test_email_ranking(self):
        emails = {"zz@gmail.com", "info@example.com", "jane@example.com", "careers@example.com"}
        ranked = rank_emails(emails, "https://example.com", ["Jane Smith"])
        self.assertEqual(ranked[:2], ["jane@example.com", "info@example.com"])
        self.assertEqual(ranked[-1], "zz@gmail.com")


class LlmParseTests(unittest.TestCase):
    def test_only_names_present_in_text_are_kept(self):
        reply = 'Sure: [{"name": "Jane Smith", "title": "CEO"}, {"name": "Made Up", "title": "Owner"}]'
        self.assertEqual(parse_people(reply, "Jane Smith is our CEO."), ["Jane Smith (CEO)"])
        self.assertEqual(parse_people("no json here", "text"), [])


class LeadsStorageTests(unittest.TestCase):
    def test_relisting_keeps_contacts_and_export_writes_them(self):
        with tempfile.TemporaryDirectory() as folder:
            db = LeadsDatabase(os.path.join(folder, "leads.sqlite3"))
            try:
                row = extract_company_rows(CARD, "https://clutch.co/x", 1)[0]
                row["website"] = "https://example.com"
                db.upsert_listing_rows([row])
                key = db.pending_companies()[0]["company_key"]
                db.save_contacts(key, {
                    "emails": ["jane@example.com", "info@example.com"], "phones": ["+1 555 123 4567"],
                    "decision_makers": ["Jane Smith (Founder & CEO)"], "linkedin": [],
                    "contact_form_url": "https://example.com/contact",
                    "contact_source_urls": ["https://example.com"], "enrichment_status": "completed",
                    "enrichment_error": "",
                })
                db.upsert_listing_rows([row])  # same company seen again on a later run
                companies = db.all_companies()
                self.assertEqual(len(companies), 1)
                self.assertEqual(companies[0]["emails"], "jane@example.com; info@example.com")
                self.assertEqual(db.pending_companies(), [])

                path = write_tracker(companies, str(Path(folder) / "tracker.xlsx"))
                wb = load_workbook(path)
                try:
                    ws = wb["Agency Tracker"]
                    header = [cell.value for cell in ws[2]]
                    values = dict(zip(header, [cell.value for cell in ws[3]]))
                finally:
                    wb.close()
            finally:
                db.close()
        self.assertEqual(values["Agency Name"], "Example AI")
        self.assertEqual(values["DM Name"], "Jane Smith")
        self.assertEqual(values["DM Title"], "Founder & CEO")
        self.assertEqual(values["Email"], "jane@example.com")
        self.assertEqual(values["Phone"], "+1 555 123 4567")
        self.assertEqual(values["Contact Form Link"], "https://example.com/contact")


class CsvTests(unittest.TestCase):
    def test_missing_fields_are_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            path = write_csv([{"agency_name": "No Site Co", "website": None}], str(Path(folder) / "out.csv"))
            with open(path, encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
        self.assertEqual(rows[0]["agency_name"], "No Site Co")
        self.assertEqual(rows[0]["website"], "")
        self.assertEqual(rows[0]["emails"], "")


class ApiFallbackTests(unittest.TestCase):
    PAGE = "<html><body><a href='mailto:info@acme-ai.com'>mail</a></body></html>"

    def test_off_by_default(self):
        finder = WebsiteContactFinder(delay=0)
        with mock.patch.object(finder, "_fetch_direct", return_value=(None, "u", "blocked:http_403")),                 mock.patch.object(website_contacts, "fetch_html") as api:
            self.assertEqual(finder.fetch("https://acme-ai.com")[2], "blocked:http_403")
            api.assert_not_called()

    def test_blocked_site_is_fetched_via_api_then_stays_on_api(self):
        finder = WebsiteContactFinder(delay=0, api_fallback=True)
        with mock.patch.object(finder, "_fetch_direct", return_value=(None, "u", "blocked:http_403")) as direct,                 mock.patch.object(website_contacts, "fetch_html", return_value=self.PAGE):
            soup, _, status = finder.fetch("https://acme-ai.com")
            finder.fetch("https://acme-ai.com/contact")
        self.assertEqual(status, "ok")
        self.assertIsNotNone(soup)
        self.assertEqual(direct.call_count, 1)  # second page went straight to the API
        self.assertEqual(finder.api_requests, 2)

    def test_robots_disallow_is_never_retried(self):
        finder = WebsiteContactFinder(delay=0, api_fallback=True)
        with mock.patch.object(finder, "_fetch_direct", return_value=(None, "u", "blocked:robots_disallow")),                 mock.patch.object(website_contacts, "fetch_html") as api:
            self.assertEqual(finder.fetch("https://acme-ai.com")[2], "blocked:robots_disallow")
            api.assert_not_called()

    def test_api_block_is_reported_not_retried(self):
        finder = WebsiteContactFinder(delay=0, api_fallback=True)
        with mock.patch.object(finder, "_fetch_direct", return_value=(None, "u", "blocked:challenge")),                 mock.patch.object(website_contacts, "fetch_html", side_effect=AccessBlockedError("x")):
            self.assertEqual(finder.fetch("https://acme-ai.com")[2], "blocked:api_challenge")


if __name__ == "__main__":
    unittest.main()
